from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import logging
import uuid
import bcrypt
import requests
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional
from datetime import datetime, timezone, timedelta
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

EMERGENT_SESSION_URL = "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data"


def now_utc():
    return datetime.now(timezone.utc)


def clean(doc):
    if doc and "_id" in doc:
        doc.pop("_id", None)
    return doc


def client_tipe(penyewa_id: str) -> str:
    if not penyewa_id:
        return "P"
    p = penyewa_id[0].upper()
    return p if p in ("A", "D", "P") else "P"


TIPE_LABEL = {"A": "Agen", "D": "Kedinasan", "P": "Perorangan"}


class OrderDriverItem(BaseModel):
    driver_id: Optional[str] = None
    driver_nama: str
    gaji: float = 0
    segmen: Optional[str] = ""
    tanggal: Optional[str] = None


class OrderCreate(BaseModel):
    id_order: Optional[str] = None
    penyewa_id: Optional[str] = None
    penyewa_nama: str
    tamu: Optional[str] = ""
    unit_id: Optional[str] = None
    unit_nama: str
    tanggal_mulai: Optional[str] = None
    tanggal_selesai: Optional[str] = None
    rute: Optional[str] = ""
    harga: float = 0
    biaya_sewa_rekanan: float = 0
    biaya_bbm: float = 0
    biaya_toll_parkir: float = 0
    biaya_lain: float = 0
    ket: Optional[str] = ""
    status_bayar: Optional[str] = "Belum Bayar"
    nominal_dp: Optional[float] = 0
    drivers: List[OrderDriverItem] = []


class ClientModel(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    penyewa_id: str
    nama: str
    tipe: Optional[str] = None


class UnitModel(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    unit_id: str
    nama: str


class DriverModel(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    driver_id: Optional[str] = None
    nama: str
    telepon: Optional[str] = ""
    aktif: bool = True
    foto_sim: Optional[str] = None


async def get_current_user(request: Request):
    token = request.cookies.get("session_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    session = await db.user_sessions.find_one({"session_token": token})
    if not session:
        raise HTTPException(status_code=401, detail="Invalid session")
    expires_at = session["expires_at"]
    if isinstance(expires_at, str):
        expires_at = datetime.fromisoformat(expires_at)
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    if expires_at < now_utc():
        raise HTTPException(status_code=401, detail="Session expired")
    user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan. Hubungi pemilik.")
    return user


def require_roles(*roles):
    async def checker(user=Depends(get_current_user)):
        if user.get("role") not in roles:
            raise HTTPException(status_code=403, detail="Akses ditolak untuk peran ini")
        return user
    return checker


@api_router.post("/auth/session")
async def auth_session(request: Request, response: Response):
    body = await request.json()
    session_id = body.get("session_id")
    if not session_id:
        raise HTTPException(status_code=400, detail="session_id required")
    r = requests.get(EMERGENT_SESSION_URL, headers={"X-Session-ID": session_id}, timeout=15)
    if r.status_code != 200:
        raise HTTPException(status_code=401, detail="Gagal memverifikasi sesi Google")
    data = r.json()
    email = data["email"]
    existing = await db.users.find_one({"email": email}, {"_id": 0})
    if existing:
        user_id = existing["user_id"]
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": data.get("name"), "picture": data.get("picture")}})
        user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    else:
        count = await db.users.count_documents({})
        role = "owner" if count == 0 else "operator"
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        user = {
            "user_id": user_id,
            "email": email,
            "name": data.get("name"),
            "picture": data.get("picture"),
            "role": role,
            "active": True,
            "created_at": now_utc().isoformat(),
        }
        await db.users.insert_one(dict(user))
        user = clean(user)

    session_token = data["session_token"]
    expires = now_utc() + timedelta(days=7)
    await db.user_sessions.insert_one({
        "user_id": user["user_id"],
        "session_token": session_token,
        "expires_at": expires.isoformat(),
        "created_at": now_utc().isoformat(),
    })
    response.set_cookie("session_token", session_token, httponly=True, secure=True,
                        samesite="none", path="/", max_age=7 * 24 * 3600)
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    return {"user": user}


@api_router.get("/auth/me")
async def auth_me(user=Depends(get_current_user)):
    return user


@api_router.post("/auth/logout")
async def auth_logout(request: Request, response: Response):
    token = request.cookies.get("session_token")
    if token:
        await db.user_sessions.delete_many({"session_token": token})
    response.delete_cookie("session_token", path="/")
    return {"ok": True}


def hash_password(p: str) -> str:
    return bcrypt.hashpw(p.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(p: str, h: str) -> bool:
    try:
        return bcrypt.checkpw(p.encode("utf-8"), h.encode("utf-8"))
    except Exception:
        return False


class LoginBody(BaseModel):
    email: str
    password: str


@api_router.post("/auth/login")
async def auth_login(body: LoginBody, request: Request, response: Response):
    email = body.email.strip().lower()
    ip = request.client.host if request.client else "?"
    ident = f"{ip}:{email}"
    att = await db.login_attempts.find_one({"identifier": ident})
    if att and att.get("count", 0) >= 5 and att.get("locked_until"):
        lu = att["locked_until"]
        lu = datetime.fromisoformat(lu) if isinstance(lu, str) else lu
        if lu.tzinfo is None:
            lu = lu.replace(tzinfo=timezone.utc)
        if lu > now_utc():
            raise HTTPException(status_code=429, detail="Terlalu banyak percobaan. Coba lagi dalam 15 menit.")
    user = await db.users.find_one({"email": email})
    if not user or not user.get("password_hash") or not verify_password(body.password, user["password_hash"]):
        cnt = (att.get("count", 0) + 1) if att else 1
        upd = {"identifier": ident, "count": cnt}
        if cnt >= 5:
            upd["locked_until"] = (now_utc() + timedelta(minutes=15)).isoformat()
        await db.login_attempts.update_one({"identifier": ident}, {"$set": upd}, upsert=True)
        raise HTTPException(status_code=401, detail="Email atau password salah")
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan. Hubungi pemilik.")
    await db.login_attempts.delete_one({"identifier": ident})
    token = f"sess_{uuid.uuid4().hex}"
    await db.user_sessions.insert_one({
        "user_id": user["user_id"], "session_token": token,
        "expires_at": (now_utc() + timedelta(days=7)).isoformat(), "created_at": now_utc().isoformat(),
    })
    response.set_cookie("session_token", token, httponly=True, secure=True, samesite="none", path="/", max_age=7 * 24 * 3600)
    clean(user)
    user.pop("password_hash", None)
    return {"user": user}


class ResetPasswordBody(BaseModel):
    email: str
    new_password: str


@api_router.post("/auth/reset-password")
async def auth_reset_password(body: ResetPasswordBody):
    email = body.email.strip().lower()
    if len(body.new_password) < 6:
        raise HTTPException(status_code=400, detail="Password minimal 6 karakter")
    user = await db.users.find_one({"email": email})
    if not user:
        raise HTTPException(status_code=404, detail="Email tidak terdaftar")
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan. Hubungi pemilik.")
    await db.users.update_one({"email": email}, {"$set": {"password_hash": hash_password(body.new_password)}})
    await db.login_attempts.delete_many({"identifier": {"$regex": f":{email}$"}})
    return {"ok": True, "message": "Password berhasil diperbarui. Silakan masuk."}


@api_router.get("/users")
async def list_users(user=Depends(require_roles("owner", "admin", "operator"))):
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)


class CreateUserBody(BaseModel):
    email: str
    password: str
    name: Optional[str] = ""
    role: str = "operator"


@api_router.post("/users")
async def create_user_account(body: CreateUserBody, user=Depends(require_roles("owner", "operator"))):
    email = body.email.strip().lower()
    if body.role not in ("owner", "admin", "operator"):
        raise HTTPException(status_code=400, detail="Peran tidak valid")
    if len(body.password) < 6:
        raise HTTPException(status_code=400, detail="Password minimal 6 karakter")
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {"user_id": f"user_{uuid.uuid4().hex[:12]}", "email": email,
           "name": body.name or email.split("@")[0], "role": body.role, "active": True,
           "password_hash": hash_password(body.password), "created_at": now_utc().isoformat()}
    await db.users.insert_one(dict(doc))
    doc.pop("password_hash", None)
    clean(doc)
    return doc


@api_router.put("/users/{user_id}")
async def update_user(user_id: str, body: dict, user=Depends(get_current_user)):
    role = user.get("role")
    is_self = user_id == user.get("user_id")
    # owner & operator: kelola semua pengguna. admin: hanya boleh ubah akunnya sendiri.
    if role not in ("owner", "operator") and not (role == "admin" and is_self):
        raise HTTPException(status_code=403, detail="Anda tidak memiliki akses untuk mengubah pengguna ini")
    target = await db.users.find_one({"user_id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    # admin yang mengubah dirinya sendiri hanya boleh ubah nama/email/password
    if role == "admin":
        body = {k: v for k, v in body.items() if k in ("name", "email", "password")}
    upd = {}
    if "role" in body and body["role"] in ("owner", "admin", "operator"):
        upd["role"] = body["role"]
    if "active" in body:
        upd["active"] = bool(body["active"])
    if "name" in body and body["name"]:
        upd["name"] = body["name"]
    if "email" in body and body["email"]:
        new_email = body["email"].strip().lower()
        if new_email != target.get("email"):
            if await db.users.find_one({"email": new_email, "user_id": {"$ne": user_id}}):
                raise HTTPException(status_code=400, detail="Email sudah terdaftar")
            upd["email"] = new_email
    if body.get("password"):
        if len(body["password"]) < 6:
            raise HTTPException(status_code=400, detail="Password minimal 6 karakter")
        upd["password_hash"] = hash_password(body["password"])
    if upd:
        await db.users.update_one({"user_id": user_id}, {"$set": upd})
    return await db.users.find_one({"user_id": user_id}, {"_id": 0, "password_hash": 0})


@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, user=Depends(require_roles("owner", "operator"))):
    if user_id == user["user_id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus akun sendiri")
    target = await db.users.find_one({"user_id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    if target.get("role") == "owner":
        owners = await db.users.count_documents({"role": "owner"})
        if owners <= 1:
            raise HTTPException(status_code=400, detail="Tidak dapat menghapus satu-satunya Pemilik")
    await db.users.delete_one({"user_id": user_id})
    await db.user_sessions.delete_many({"user_id": user_id})
    return {"ok": True}


@api_router.get("/clients")
async def list_clients(user=Depends(get_current_user)):
    return await db.clients.find({}, {"_id": 0}).sort("penyewa_id", 1).to_list(5000)


@api_router.post("/clients")
async def create_client(body: ClientModel, user=Depends(require_roles("owner", "admin", "operator"))):
    body.tipe = client_tipe(body.penyewa_id)
    doc = body.model_dump()
    await db.clients.insert_one(dict(doc))
    return clean(doc)


@api_router.put("/clients/{id}")
async def update_client(id: str, body: ClientModel, user=Depends(require_roles("owner", "admin", "operator"))):
    existing = await db.clients.find_one({"id": id})
    body.tipe = client_tipe(body.penyewa_id)
    doc = body.model_dump()
    doc["id"] = id
    await db.clients.update_one({"id": id}, {"$set": doc})
    old_pid = (existing or {}).get("penyewa_id")
    if old_pid:
        await db.orders.update_many({"penyewa_id": old_pid}, {"$set": {
            "penyewa_nama": doc["nama"], "penyewa_id": doc["penyewa_id"], "penyewa_tipe": doc["tipe"]}})
    return clean(doc)


@api_router.delete("/clients/{id}")
async def delete_client(id: str, user=Depends(require_roles("owner", "admin"))):
    await db.clients.delete_one({"id": id})
    return {"ok": True}


@api_router.get("/units")
async def list_units(user=Depends(get_current_user)):
    return await db.units.find({}, {"_id": 0}).sort("unit_id", 1).to_list(5000)


@api_router.post("/units")
async def create_unit(body: UnitModel, user=Depends(require_roles("owner", "admin", "operator"))):
    doc = body.model_dump()
    await db.units.insert_one(dict(doc))
    return clean(doc)


@api_router.put("/units/{id}")
async def update_unit(id: str, body: UnitModel, user=Depends(require_roles("owner", "admin", "operator"))):
    existing = await db.units.find_one({"id": id})
    doc = body.model_dump()
    doc["id"] = id
    await db.units.update_one({"id": id}, {"$set": doc})
    old_uid = (existing or {}).get("unit_id")
    if old_uid:
        await db.orders.update_many({"unit_id": old_uid}, {"$set": {
            "unit_nama": doc["nama"], "unit_id": doc["unit_id"]}})
    return clean(doc)


@api_router.delete("/units/{id}")
async def delete_unit(id: str, user=Depends(require_roles("owner", "admin"))):
    await db.units.delete_one({"id": id})
    return {"ok": True}


@api_router.get("/drivers")
async def list_drivers(user=Depends(get_current_user)):
    return await db.drivers.find({}, {"_id": 0}).sort("nama", 1).to_list(5000)


async def gen_driver_code():
    drivers = await db.drivers.find({}, {"driver_id": 1}).to_list(10000)
    maxn = 0
    for d in drivers:
        did = d.get("driver_id") or ""
        if did.startswith("DR") and did[2:].isdigit():
            maxn = max(maxn, int(did[2:]))
    return f"DR{maxn + 1:03d}"


@api_router.get("/drivers/next-code")
async def drivers_next_code(user=Depends(get_current_user)):
    return {"driver_id": await gen_driver_code()}


@api_router.post("/drivers")
async def create_driver(body: DriverModel, user=Depends(require_roles("owner", "admin", "operator"))):
    doc = body.model_dump()
    if not doc.get("driver_id"):
        doc["driver_id"] = await gen_driver_code()
    await db.drivers.insert_one(dict(doc))
    return clean(doc)


@api_router.put("/drivers/{id}")
async def update_driver(id: str, body: DriverModel, user=Depends(require_roles("owner", "admin", "operator"))):
    existing = await db.drivers.find_one({"id": id})
    doc = body.model_dump()
    doc["id"] = id
    if not doc.get("driver_id"):
        doc["driver_id"] = (existing or {}).get("driver_id") or await gen_driver_code()
    await db.drivers.update_one({"id": id}, {"$set": doc})
    old_code = (existing or {}).get("driver_id")
    if old_code:
        await db.orders.update_many(
            {"drivers.driver_id": old_code},
            {"$set": {"drivers.$[e].driver_nama": doc["nama"], "drivers.$[e].driver_id": doc["driver_id"]}},
            array_filters=[{"e.driver_id": old_code}],
        )
        await db.kasbon.update_many({"driver_id": old_code}, {"$set": {"driver_nama": doc["nama"], "driver_id": doc["driver_id"]}})
    return clean(doc)


@api_router.delete("/drivers/{id}")
async def delete_driver(id: str, user=Depends(require_roles("owner", "admin"))):
    await db.drivers.delete_one({"id": id})
    return {"ok": True}


def is_lengkap(o: dict) -> bool:
    basic = all([o.get("penyewa_nama"), o.get("tamu"), o.get("unit_nama"),
                 o.get("tanggal_mulai"), o.get("tanggal_selesai"), o.get("rute")]) and float(o.get("harga") or 0) > 0
    drivers = o.get("drivers") or []
    drv_ok = len(drivers) > 0 and all((d.get("driver_nama") and float(d.get("gaji") or 0) > 0) for d in drivers)
    return bool(basic and drv_ok)


def compute_order(o: dict) -> dict:
    total_gaji = sum(float(d.get("gaji") or 0) for d in o.get("drivers", []))
    total_biaya = (float(o.get("biaya_sewa_rekanan") or 0) + total_gaji +
                   float(o.get("biaya_bbm") or 0) + float(o.get("biaya_toll_parkir") or 0) +
                   float(o.get("biaya_lain") or 0))
    o["total_gaji_driver"] = total_gaji
    o["total_biaya"] = total_biaya
    o["margin"] = float(o.get("harga") or 0) - total_biaya
    o["status"] = "lengkap" if is_lengkap(o) else "belum_lengkap"
    return o


async def gen_order_code(tanggal_mulai=None):
    dt = None
    if tanggal_mulai:
        try:
            dt = datetime.fromisoformat(tanggal_mulai)
        except Exception:
            dt = None
    if dt is None:
        dt = now_utc()
    prefix = dt.strftime("%y%m")
    existing = await db.orders.find({"id_order": {"$regex": f"^{prefix}"}}, {"id_order": 1}).to_list(20000)
    maxseq = 0
    for e in existing:
        tail = (e.get("id_order") or "")[len(prefix):]
        if tail.isdigit():
            maxseq = max(maxseq, int(tail))
    return f"{prefix}{maxseq + 1:05d}"


@api_router.get("/orders/next-code")
async def orders_next_code(user=Depends(get_current_user), tanggal: Optional[str] = None):
    return {"id_order": await gen_order_code(tanggal)}


@api_router.get("/orders")
async def list_orders(user=Depends(get_current_user), search: Optional[str] = None,
                      start: Optional[str] = None, end: Optional[str] = None, status: Optional[str] = None,
                      bulan: Optional[int] = None, tahun: Optional[int] = None):
    q = {}
    if search:
        q["$or"] = [{"penyewa_nama": {"$regex": search, "$options": "i"}},
                    {"tamu": {"$regex": search, "$options": "i"}},
                    {"rute": {"$regex": search, "$options": "i"}},
                    {"id_order": {"$regex": search, "$options": "i"}}]
    if start:
        q.setdefault("tanggal_mulai", {})["$gte"] = start
    if end:
        q.setdefault("tanggal_mulai", {})["$lte"] = end
    if status in ("lengkap", "belum_lengkap"):
        q["status"] = status
    by_period = bulan and tahun
    if by_period:
        q["tanggal_mulai"] = {"$regex": f"^{int(tahun):04d}-{int(bulan):02d}"}
    sort_field, sort_dir = ("id_order", 1) if by_period else ("tanggal_mulai", -1)
    return await db.orders.find(q, {"_id": 0}).sort(sort_field, sort_dir).to_list(5000)


@api_router.post("/orders")
async def create_order(body: OrderCreate, user=Depends(require_roles("owner", "admin", "operator"))):
    o = body.model_dump()
    total_gaji = sum(float(d.get("gaji") or 0) for d in o.get("drivers", []))
    if total_gaji > float(o.get("harga") or 0):
        raise HTTPException(status_code=400, detail="Total gaji driver tidak boleh melebihi harga sewa")
    o["id"] = str(uuid.uuid4())
    provided = (o.get("id_order") or "").strip()
    if provided:
        if await db.orders.find_one({"id_order": provided}):
            raise HTTPException(status_code=400, detail=f"ID Order {provided} sudah digunakan")
        o["id_order"] = provided
    else:
        o["id_order"] = await gen_order_code(o.get("tanggal_mulai"))
    o["penyewa_tipe"] = client_tipe(o.get("penyewa_id") or "")
    compute_order(o)
    o["created_by"] = user["email"]
    o["created_at"] = now_utc().isoformat()
    await db.orders.insert_one(dict(o))
    return clean(o)


@api_router.get("/orders/{id}")
async def get_order(id: str, user=Depends(get_current_user)):
    o = await db.orders.find_one({"id": id}, {"_id": 0})
    if not o:
        raise HTTPException(status_code=404, detail="Order tidak ditemukan")
    return o


@api_router.put("/orders/{id}")
async def update_order(id: str, body: OrderCreate, user=Depends(require_roles("owner", "admin", "operator"))):
    existing = await db.orders.find_one({"id": id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Order tidak ditemukan")
    o = body.model_dump()
    total_gaji = sum(float(d.get("gaji") or 0) for d in o.get("drivers", []))
    if total_gaji > float(o.get("harga") or 0):
        raise HTTPException(status_code=400, detail="Total gaji driver tidak boleh melebihi harga sewa")
    o["id"] = id
    provided = (o.get("id_order") or "").strip()
    if provided and provided != existing.get("id_order"):
        if await db.orders.find_one({"id_order": provided, "id": {"$ne": id}}):
            raise HTTPException(status_code=400, detail=f"ID Order {provided} sudah digunakan")
        o["id_order"] = provided
    else:
        o["id_order"] = existing.get("id_order")
    o["penyewa_tipe"] = client_tipe(o.get("penyewa_id") or "")
    compute_order(o)
    o["created_by"] = existing.get("created_by")
    o["created_at"] = existing.get("created_at")
    o["updated_at"] = now_utc().isoformat()
    await db.orders.update_one({"id": id}, {"$set": o})
    return clean(o)


@api_router.delete("/orders/{id}")
async def delete_order(id: str, user=Depends(require_roles("owner", "admin"))):
    await db.orders.delete_one({"id": id})
    return {"ok": True}


async def fetch_orders(start=None, end=None):
    q = {}
    if start:
        q.setdefault("tanggal_mulai", {})["$gte"] = start
    if end:
        q.setdefault("tanggal_mulai", {})["$lte"] = end
    return await db.orders.find(q, {"_id": 0}).to_list(10000)


def month_key(o):
    tm = o.get("tanggal_mulai")
    if not tm:
        return "N/A"
    try:
        return datetime.fromisoformat(tm).strftime("%Y-%m")
    except Exception:
        return "N/A"


@api_router.get("/analytics/executive")
async def analytics_executive(user=Depends(get_current_user), start: Optional[str] = None, end: Optional[str] = None):
    orders = await fetch_orders(start, end)
    total_omset = sum(float(o.get("harga") or 0) for o in orders)
    total_margin = sum(float(o.get("margin") or 0) for o in orders)
    units = set(o.get("unit_nama") for o in orders if o.get("unit_nama"))
    monthly = defaultdict(lambda: {"omset": 0, "margin": 0, "order": 0})
    for o in orders:
        k = month_key(o)
        monthly[k]["omset"] += float(o.get("harga") or 0)
        monthly[k]["margin"] += float(o.get("margin") or 0)
        monthly[k]["order"] += 1
    trend = [{"bulan": k, **v} for k, v in sorted(monthly.items()) if k != "N/A"]
    return {"total_omset": total_omset, "total_margin": total_margin,
            "total_order": len(orders), "total_unit": len(units), "trend": trend}


@api_router.get("/analytics/clients")
async def analytics_clients(user=Depends(get_current_user), start: Optional[str] = None,
                            end: Optional[str] = None, tipe: Optional[str] = None):
    orders = await fetch_orders(start, end)
    if tipe:
        orders = [o for o in orders if o.get("penyewa_tipe") == tipe]
    by_tipe = defaultdict(lambda: {"order": 0, "omset": 0, "margin": 0})
    by_client = defaultdict(lambda: {"order": 0, "omset": 0, "margin": 0})
    for o in orders:
        t = o.get("penyewa_tipe") or "P"
        by_tipe[t]["order"] += 1
        by_tipe[t]["omset"] += float(o.get("harga") or 0)
        by_tipe[t]["margin"] += float(o.get("margin") or 0)
        c = o.get("penyewa_nama") or "-"
        by_client[c]["order"] += 1
        by_client[c]["omset"] += float(o.get("harga") or 0)
        by_client[c]["margin"] += float(o.get("margin") or 0)
    tipe_data = [{"tipe": k, "label": TIPE_LABEL.get(k, k), **v} for k, v in by_tipe.items()]
    top_clients = sorted([{"nama": k, **v} for k, v in by_client.items()], key=lambda x: x["omset"], reverse=True)
    return {"by_tipe": tipe_data, "top_clients": top_clients}


@api_router.get("/analytics/units")
async def analytics_units(user=Depends(get_current_user), start: Optional[str] = None,
                          end: Optional[str] = None, unit: Optional[str] = None):
    orders = await fetch_orders(start, end)
    if unit:
        orders = [o for o in orders if (o.get("unit_nama") or "-") == unit]
    by_unit = defaultdict(lambda: {"order": 0, "omset": 0, "margin": 0})
    monthly = defaultdict(lambda: defaultdict(int))
    for o in orders:
        u = o.get("unit_nama") or "-"
        by_unit[u]["order"] += 1
        by_unit[u]["omset"] += float(o.get("harga") or 0)
        by_unit[u]["margin"] += float(o.get("margin") or 0)
        monthly[month_key(o)][u] += 1
    summary = []
    for k, v in by_unit.items():
        avg = v["omset"] / v["order"] if v["order"] else 0
        summary.append({"unit": k, **v, "rata_omset": avg})
    summary.sort(key=lambda x: x["omset"], reverse=True)
    LAINNYA = "Lainnya"
    freq = []
    key_totals = defaultdict(float)
    has_lainnya = False
    for k in sorted(monthly.keys()):
        if k == "N/A":
            continue
        ranked = sorted(monthly[k].items(), key=lambda x: x[1], reverse=True)
        row = {"bulan": k}
        for name, c in ranked[:5]:
            row[name] = c
            key_totals[name] += c
        rest = sum(c for _, c in ranked[5:])
        if rest > 0:
            row[LAINNYA] = rest
            has_lainnya = True
        freq.append(row)
    freq_keys = [n for n, _ in sorted(key_totals.items(), key=lambda x: x[1], reverse=True)]
    if has_lainnya:
        freq_keys.append(LAINNYA)
    return {"summary": summary, "freq": freq, "units": [s["unit"] for s in summary], "freq_keys": freq_keys}


@api_router.get("/analytics/drivers")
async def analytics_drivers(user=Depends(get_current_user), bulan: Optional[int] = None, tahun: Optional[int] = None):
    orders = await db.orders.find({}, {"_id": 0}).to_list(10000)
    by_driver = defaultdict(lambda: {"tugas": 0, "total_gaji": 0, "tasks": []})
    for o in orders:
        mk = month_key(o)
        if mk != "N/A":
            y, m = mk.split("-")
            if tahun and int(y) != tahun:
                continue
            if bulan and int(m) != bulan:
                continue
        elif tahun or bulan:
            continue
        for d in o.get("drivers", []):
            nama = d.get("driver_nama") or "-"
            by_driver[nama]["tugas"] += 1
            by_driver[nama]["total_gaji"] += float(d.get("gaji") or 0)
            by_driver[nama]["tasks"].append({
                "id_order": o.get("id_order"),
                "penyewa_nama": o.get("penyewa_nama"),
                "tanggal_mulai": o.get("tanggal_mulai"),
                "tanggal_selesai": o.get("tanggal_selesai"),
                "rute": o.get("rute"),
                "unit_nama": o.get("unit_nama"),
                "segmen": d.get("segmen") or "",
                "gaji": float(d.get("gaji") or 0),
            })
    for v in by_driver.values():
        v["tasks"].sort(key=lambda t: t.get("id_order") or "")
    data = sorted([{"nama": k, **v} for k, v in by_driver.items()], key=lambda x: x["tugas"], reverse=True)
    return {"drivers": data}


def make_xlsx(title, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:30]
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    for ri, row in enumerate(rows, 2):
        for ci, val in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=val)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 45)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_response(buf, filename):
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@api_router.get("/export/orders")
async def export_orders(user=Depends(get_current_user), search: Optional[str] = None,
                        start: Optional[str] = None, end: Optional[str] = None, status: Optional[str] = None,
                        bulan: Optional[int] = None, tahun: Optional[int] = None):
    orders = await list_orders(user, search, start, end, status, bulan, tahun)
    status_label = {"lengkap": "Lengkap", "belum_lengkap": "Belum Lengkap"}
    headers = ["ID Order", "Penyewa", "Tipe", "Tamu", "Unit", "Tgl Mulai", "Tgl Selesai", "Rute",
               "Harga", "Sewa Mobil", "Gaji Driver", "BBM", "Tol/Parkir", "Lain-lain",
               "Total Biaya", "Margin", "Driver", "Status Kelengkapan", "Status Pembayaran", "Nominal DP"]
    rows = []
    for o in orders:
        drv = "; ".join(f"{d.get('driver_nama')} ({d.get('segmen','')}) Rp{int(d.get('gaji') or 0)}" for d in o.get("drivers", []))
        rows.append([o.get("id_order"), o.get("penyewa_nama"), TIPE_LABEL.get(o.get("penyewa_tipe"), ""),
                     o.get("tamu"), o.get("unit_nama"), o.get("tanggal_mulai"), o.get("tanggal_selesai"),
                     o.get("rute"), o.get("harga"), o.get("biaya_sewa_rekanan"), o.get("total_gaji_driver"),
                     o.get("biaya_bbm"), o.get("biaya_toll_parkir"), o.get("biaya_lain"),
                     o.get("total_biaya"), o.get("margin"), drv,
                     status_label.get(o.get("status"), o.get("status")), o.get("status_bayar") or "Belum Bayar",
                     o.get("nominal_dp") or 0])
    return xlsx_response(make_xlsx("Transaksi", headers, rows), "transaksi.xlsx")


@api_router.get("/export/clients-analysis")
async def export_clients(user=Depends(get_current_user), start: Optional[str] = None,
                         end: Optional[str] = None, tipe: Optional[str] = None):
    data = await analytics_clients(user, start, end, tipe)
    headers = ["Penyewa", "Jumlah Order", "Omset", "Margin"]
    rows = [[c["nama"], c["order"], c["omset"], c["margin"]] for c in data["top_clients"]]
    return xlsx_response(make_xlsx("Analisis Klien", headers, rows), "analisis_klien.xlsx")


@api_router.get("/export/units-analysis")
async def export_units(user=Depends(get_current_user), start: Optional[str] = None,
                       end: Optional[str] = None, unit: Optional[str] = None):
    data = await analytics_units(user, start, end, unit)
    headers = ["Unit", "Jumlah Order", "Total Omset", "Rata-rata Omset", "Margin"]
    rows = [[u["unit"], u["order"], u["omset"], round(u["rata_omset"]), u["margin"]] for u in data["summary"]]
    return xlsx_response(make_xlsx("Analisis Unit", headers, rows), "analisis_unit.xlsx")


@api_router.get("/export/drivers-analysis")
async def export_drivers(user=Depends(get_current_user), bulan: Optional[int] = None, tahun: Optional[int] = None):
    data = await analytics_drivers(user, bulan, tahun)
    headers = ["Driver", "Jumlah Tugas", "Total Gaji"]
    rows = [[d["nama"], d["tugas"], d["total_gaji"]] for d in data["drivers"]]
    return xlsx_response(make_xlsx("Analisis Driver", headers, rows), "analisis_driver.xlsx")


# ============================ KASBON DRIVER ============================
class KasbonCreate(BaseModel):
    driver_id: str
    tanggal: str
    jumlah: float


@api_router.post("/kasbon")
async def create_kasbon(body: KasbonCreate, user=Depends(require_roles("owner", "admin", "operator"))):
    drv = await db.drivers.find_one({"driver_id": body.driver_id}, {"_id": 0})
    doc = {"id": str(uuid.uuid4()), "driver_id": body.driver_id,
           "driver_nama": (drv or {}).get("nama") or body.driver_id,
           "tanggal": body.tanggal, "jumlah": float(body.jumlah or 0), "jenis": "pinjam",
           "created_by": user["email"], "created_at": now_utc().isoformat()}
    await db.kasbon.insert_one(dict(doc))
    return clean(doc)


@api_router.get("/kasbon/summary")
async def kasbon_summary(user=Depends(get_current_user)):
    rows = await db.kasbon.find({}, {"_id": 0}).to_list(50000)
    agg = defaultdict(lambda: {"jumlah": 0, "total_pinjam": 0.0, "total_potong": 0.0, "driver_nama": ""})
    for r in rows:
        did = r.get("driver_id") or "-"
        a = agg[did]
        a["driver_nama"] = r.get("driver_nama") or did
        if r.get("jenis") == "potong":
            a["total_potong"] += float(r.get("jumlah") or 0)
        else:
            a["jumlah"] += 1
            a["total_pinjam"] += float(r.get("jumlah") or 0)
    out = [{"driver_id": did, "driver_nama": a["driver_nama"], "jumlah": a["jumlah"],
            "total_pinjam": a["total_pinjam"], "total_potong": a["total_potong"],
            "sisa": a["total_pinjam"] - a["total_potong"]} for did, a in agg.items()]
    out.sort(key=lambda x: x["sisa"], reverse=True)
    return {"drivers": out}


@api_router.get("/kasbon/driver/{driver_id}")
async def kasbon_by_driver(driver_id: str, user=Depends(get_current_user)):
    rows = await db.kasbon.find({"driver_id": driver_id}, {"_id": 0}).to_list(50000)
    pinjam = sorted([r for r in rows if r.get("jenis") != "potong"], key=lambda x: x.get("tanggal") or "")
    potong = sorted([r for r in rows if r.get("jenis") == "potong"], key=lambda x: x.get("tanggal") or "")
    total_pinjam = sum(float(r.get("jumlah") or 0) for r in pinjam)
    total_potong = sum(float(r.get("jumlah") or 0) for r in potong)
    drv = await db.drivers.find_one({"driver_id": driver_id}, {"_id": 0})
    return {"driver_id": driver_id, "driver_nama": (drv or {}).get("nama") or driver_id,
            "pinjam": pinjam, "potong": potong, "total_pinjam": total_pinjam,
            "total_potong": total_potong, "sisa": total_pinjam - total_potong}


@api_router.delete("/kasbon/{id}")
async def delete_kasbon(id: str, user=Depends(require_roles("owner", "admin"))):
    await db.kasbon.delete_one({"id": id})
    return {"ok": True}


# ============================ PENGGAJIAN / SLIP GAJI ============================
BULAN_ID = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
            "Agustus", "September", "Oktober", "November", "Desember"]


def _rp(n):
    return "Rp " + format(int(round(float(n or 0))), ",d").replace(",", ".")


def payroll_range(tahun, bulan, periode):
    import calendar
    if int(periode) == 1:
        return f"{int(tahun):04d}-{int(bulan):02d}-01", f"{int(tahun):04d}-{int(bulan):02d}-15"
    last = calendar.monthrange(int(tahun), int(bulan))[1]
    return f"{int(tahun):04d}-{int(bulan):02d}-16", f"{int(tahun):04d}-{int(bulan):02d}-{last:02d}"


@api_router.get("/payroll")
async def payroll(user=Depends(get_current_user), tahun: Optional[int] = None,
                  bulan: Optional[int] = None, periode: int = 1, driver_id: Optional[str] = None):
    if not tahun or not bulan:
        return {"drivers": []}
    start, end = payroll_range(tahun, bulan, periode)
    orders = await db.orders.find({}, {"_id": 0}).to_list(50000)
    by_driver = defaultdict(lambda: {"driver_nama": "", "items": [], "total_gaji": 0.0})
    for o in orders:
        for d in o.get("drivers", []):
            tg = d.get("tanggal") or o.get("tanggal_mulai")
            if not tg or not (start <= tg <= end):
                continue
            did = d.get("driver_id") or d.get("driver_nama") or "-"
            if driver_id and did != driver_id:
                continue
            a = by_driver[did]
            a["driver_nama"] = d.get("driver_nama") or did
            a["items"].append({"tanggal": tg, "id_order": o.get("id_order"),
                               "rute": d.get("segmen") or o.get("rute"), "unit_nama": o.get("unit_nama"),
                               "penyewa_nama": o.get("penyewa_nama"), "gaji": float(d.get("gaji") or 0)})
            a["total_gaji"] += float(d.get("gaji") or 0)
    result = []
    for did, a in by_driver.items():
        a["items"].sort(key=lambda x: x["tanggal"])
        krows = await db.kasbon.find({"driver_id": did}, {"_id": 0}).to_list(20000)
        sisa = (sum(float(r.get("jumlah") or 0) for r in krows if r.get("jenis") != "potong")
                - sum(float(r.get("jumlah") or 0) for r in krows if r.get("jenis") == "potong"))
        result.append({"driver_id": did, "driver_nama": a["driver_nama"], "items": a["items"],
                       "total_gaji": a["total_gaji"], "sisa_kasbon": sisa})
    result.sort(key=lambda x: x["total_gaji"], reverse=True)
    return {"periode": {"start": start, "end": end, "tahun": tahun, "bulan": bulan, "periode": periode},
            "drivers": result}


class SlipBody(BaseModel):
    driver_id: str
    tahun: int
    bulan: int
    periode: int = 1
    potongan_kasbon: float = 0


def build_slip_pdf(drv, tahun, bulan, periode, potong, payday):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=18 * mm, rightMargin=18 * mm)
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Title"], fontSize=15, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#059669"))
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9)
    periode_lbl = "Tanggal 1 - 15" if int(periode) == 1 else "Tanggal 16 - Akhir Bulan"
    el = []
    el.append(Paragraph("PT. JAWA DWIPA SOLUTIONS", h))
    el.append(Paragraph("Slip Gaji Driver &middot; Surabaya", sub))
    el.append(Spacer(1, 10))
    info = [["Nama Driver", ": " + str(drv.get("driver_nama") or "-"), "Periode", ": " + periode_lbl],
            ["ID Driver", ": " + str(drv.get("driver_id") or "-"), "Bulan", f": {BULAN_ID[int(bulan)]} {int(tahun)}"],
            ["Tanggal Bayar", ": " + str(payday), "", ""]]
    ti = Table(info, colWidths=[75, 150, 70, 130])
    ti.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    el.append(ti)
    el.append(Spacer(1, 12))
    data = [["Tanggal", "ID Order", "Rute", "Unit", "Gaji"]]
    for it in drv.get("items", []):
        data.append([it.get("tanggal"), it.get("id_order"), (it.get("rute") or "")[:38],
                     it.get("unit_nama") or "", _rp(it.get("gaji"))])
    if len(data) == 1:
        data.append(["-", "-", "Tidak ada penugasan", "-", _rp(0)])
    t = Table(data, colWidths=[62, 70, 180, 80, 78], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#059669")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (4, 0), (4, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    el.append(t)
    el.append(Spacer(1, 10))
    total = float(drv.get("total_gaji") or 0)
    bersih = total - float(potong or 0)
    summ = [["Total Gaji Kotor", _rp(total)],
            ["Potongan Kasbon", "- " + _rp(potong)],
            ["Gaji Bersih Diterima", _rp(bersih)]]
    ts = Table(summ, colWidths=[300, 173])
    ts.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEABOVE", (0, 2), (-1, 2), 0.8, colors.HexColor("#0f172a")),
        ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 2), (-1, 2), colors.HexColor("#059669")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    el.append(ts)
    el.append(Spacer(1, 6))
    el.append(Paragraph(f"Sisa kasbon setelah potongan: {_rp(float(drv.get('sisa_kasbon') or 0) - float(potong or 0))}", small))
    el.append(Spacer(1, 30))
    sign = [["Diterima oleh,", "", "Hormat kami,"],
            ["", "", ""], ["", "", ""],
            [f"( {drv.get('driver_nama') or '________'} )", "", "( ________________ )"]]
    tsign = Table(sign, colWidths=[190, 90, 190])
    tsign.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    el.append(tsign)
    doc.build(el)
    buf.seek(0)
    return buf


@api_router.post("/payroll/slip")
async def payroll_slip(body: SlipBody, user=Depends(require_roles("owner", "admin", "operator"))):
    data = await payroll(user=user, tahun=body.tahun, bulan=body.bulan, periode=body.periode, driver_id=body.driver_id)
    drv = next((d for d in data["drivers"] if d["driver_id"] == body.driver_id), None)
    if not drv:
        raise HTTPException(status_code=404, detail="Tidak ada data gaji untuk driver pada periode ini")
    potong = float(body.potongan_kasbon or 0)
    if potong < 0:
        raise HTTPException(status_code=400, detail="Nominal potongan tidak valid")
    payday = payroll_range(body.tahun, body.bulan, body.periode)[1]
    if potong > 0:
        await db.kasbon.insert_one({"id": str(uuid.uuid4()), "driver_id": body.driver_id,
            "driver_nama": drv["driver_nama"], "tanggal": payday, "jumlah": potong, "jenis": "potong",
            "created_by": user["email"], "created_at": now_utc().isoformat(),
            "ket": f"Potongan gaji {BULAN_ID[int(body.bulan)]} {int(body.tahun)} periode {body.periode}"})
    pdf = build_slip_pdf(drv, body.tahun, body.bulan, body.periode, potong, payday)
    safe = "".join(c for c in (drv["driver_nama"] or "driver") if c.isalnum() or c in " _-").replace(" ", "_")
    fname = f"slip_gaji_{safe}_{int(body.tahun)}{int(body.bulan):02d}_p{body.periode}.pdf"
    return StreamingResponse(pdf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


app.include_router(api_router)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()


@app.on_event("startup")
async def seed_accounts():
    try:
        await db.users.create_index("email", unique=True)
    except Exception as e:
        logger.warning(f"email index: {e}")
    try:
        await db.orders.create_index("tanggal_mulai")
        await db.orders.create_index("id_order", unique=True)
    except Exception as e:
        logger.warning(f"orders index: {e}")
    pwd = os.environ.get("DEFAULT_PASSWORD", "Jds@2025")
    owner_email = os.environ.get("OWNER_EMAIL", "adityapermanarh62@gmail.com").strip().lower()
    defaults = [
        (owner_email, "owner", "Pemilik"),
        ("admin@jds.com", "admin", "Admin Operasional"),
        ("operator@jds.com", "operator", "Operator"),
    ]
    for email, role, name in defaults:
        existing = await db.users.find_one({"email": email})
        if not existing:
            await db.users.insert_one({
                "user_id": f"user_{uuid.uuid4().hex[:12]}", "email": email, "name": name,
                "role": role, "active": True, "password_hash": hash_password(pwd),
                "created_at": now_utc().isoformat(),
            })
        else:
            setd = {}
            if role == "owner" and existing.get("role") != "owner":
                setd["role"] = "owner"
            if not existing.get("password_hash"):
                setd["password_hash"] = hash_password(pwd)
            if setd:
                await db.users.update_one({"email": email}, {"$set": setd})


@app.on_event("startup")
async def migrate_driver_ids():
    drivers = await db.drivers.find({}).to_list(20000)
    maxn = 0
    for d in drivers:
        did = d.get("driver_id") or ""
        if did.startswith("DR") and did[2:].isdigit():
            maxn = max(maxn, int(did[2:]))
    name_to_code = {}
    for d in drivers:
        if not d.get("driver_id"):
            maxn += 1
            code = f"DR{maxn:03d}"
            await db.drivers.update_one({"id": d["id"]}, {"$set": {"driver_id": code}})
            name_to_code[d.get("nama")] = code
        else:
            name_to_code[d.get("nama")] = d["driver_id"]
    orders = await db.orders.find({}).to_list(50000)
    for o in orders:
        drvs = o.get("drivers") or []
        for dr in drvs:
            code = name_to_code.get(dr.get("driver_nama"))
            if code:
                dr["driver_id"] = code
        o["drivers"] = drvs
        compute_order(o)
        await db.orders.update_one({"id": o["id"]}, {"$set": {"drivers": drvs, "status": o["status"],
            "total_gaji_driver": o["total_gaji_driver"], "total_biaya": o["total_biaya"], "margin": o["margin"]}})
    await db.drivers.update_many({}, {"$unset": {"alamat": "", "tanggal_bergabung": ""}})
    await db.orders.update_many({"status_bayar": {"$exists": False}}, {"$set": {"status_bayar": "Belum Bayar"}})
